"""
Document handling module for file operations
Supports: PDF, DOCX, XLSX, PPTX, TXT export
"""

import io
import logging
from pathlib import Path
from typing import List, Union

# Optional imports with graceful degradation
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logging.warning("reportlab not installed - PDF export will be limited")

try:
    from docx import Document
    from docx.shared import Inches, Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logging.warning("python-docx not installed - DOCX export unavailable")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logging.warning("openpyxl not installed - XLSX export unavailable")

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches, Pt as PptxPt
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    logging.warning("python-pptx not installed - PPTX export unavailable")

try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False
    logging.warning("PyPDF2 not installed - PDF reading unavailable")

# Get the data directory
ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
(DATA / "uploads").mkdir(parents=True, exist_ok=True)
(DATA / "outputs").mkdir(parents=True, exist_ok=True)


def save_upload(content: bytes, filename: str) -> Path:
    """
    Save uploaded file to uploads directory

    Args:
        content: File content as bytes
        filename: Original filename

    Returns:
        Path to saved file
    """
    path = DATA / "uploads" / filename
    path.write_bytes(content)
    logging.info(f"Saved upload: {filename} ({len(content)} bytes)")
    return path


def export_pdf(text: str, filename: str) -> Path:
    """
    Export text as PDF using reportlab

    Args:
        text: Text content to export
        filename: Output filename

    Returns:
        Path to generated PDF
    """
    output_path = DATA / "outputs" / filename

    if REPORTLAB_AVAILABLE:
        try:
            # Create PDF with reportlab
            doc = SimpleDocTemplate(str(output_path), pagesize=letter)
            story = []
            styles = getSampleStyleSheet()

            # Add title
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor='#1a1a1a',
                spaceAfter=30,
            )
            story.append(Paragraph("Generated Document", title_style))
            story.append(Spacer(1, 12))

            # Add content paragraphs
            normal_style = styles['Normal']
            for para in text.split('\n'):
                if para.strip():
                    # Escape special characters for reportlab
                    para_clean = para.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(para_clean, normal_style))
                    story.append(Spacer(1, 12))

            doc.build(story)
            logging.info(f"Generated PDF: {filename}")
            return output_path

        except Exception as e:
            logging.error(f"PDF generation failed: {e}")
            # Fallback to simple text file
            output_path = DATA / "outputs" / filename.replace('.pdf', '.txt')
            output_path.write_text(text, encoding='utf-8')
            return output_path
    else:
        # Fallback: save as text file
        logging.warning("reportlab not available, saving as TXT instead")
        output_path = DATA / "outputs" / filename.replace('.pdf', '.txt')
        output_path.write_text(text, encoding='utf-8')
        return output_path


def export_docx(text: str, filename: str) -> Path:
    """
    Export text as DOCX

    Args:
        text: Text content to export
        filename: Output filename

    Returns:
        Path to generated DOCX
    """
    output_path = DATA / "outputs" / filename

    if not DOCX_AVAILABLE:
        logging.warning("python-docx not available, saving as TXT instead")
        output_path = DATA / "outputs" / filename.replace('.docx', '.txt')
        output_path.write_text(text, encoding='utf-8')
        return output_path

    try:
        doc = Document()

        # Add title
        doc.add_heading('Generated Document', 0)

        # Add content
        for para in text.split('\n'):
            if para.strip():
                p = doc.add_paragraph(para)
                # Optional: style the paragraph
                p.style = 'Normal'

        # Save
        doc.save(str(output_path))
        logging.info(f"Generated DOCX: {filename}")
        return output_path

    except Exception as e:
        logging.error(f"DOCX generation failed: {e}")
        # Fallback to text
        output_path = DATA / "outputs" / filename.replace('.docx', '.txt')
        output_path.write_text(text, encoding='utf-8')
        return output_path


def export_xlsx(rows: List[List[str]], filename: str) -> Path:
    """
    Export data as Excel XLSX

    Args:
        rows: List of rows, each row is a list of cell values
        filename: Output filename

    Returns:
        Path to generated XLSX
    """
    output_path = DATA / "outputs" / filename

    if not OPENPYXL_AVAILABLE:
        logging.warning("openpyxl not available, saving as CSV instead")
        output_path = DATA / "outputs" / filename.replace('.xlsx', '.csv')
        csv_content = '\n'.join([','.join(map(str, row)) for row in rows])
        output_path.write_text(csv_content, encoding='utf-8')
        return output_path

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"

        # Add header formatting
        header_font = Font(bold=True, size=12)
        header_alignment = Alignment(horizontal='center', vertical='center')

        # Write rows
        for row_idx, row in enumerate(rows, 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)

                # Style header row
                if row_idx == 1:
                    cell.font = header_font
                    cell.alignment = header_alignment

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save
        wb.save(str(output_path))
        logging.info(f"Generated XLSX: {filename}")
        return output_path

    except Exception as e:
        logging.error(f"XLSX generation failed: {e}")
        # Fallback to CSV
        output_path = DATA / "outputs" / filename.replace('.xlsx', '.csv')
        csv_content = '\n'.join([','.join(map(str, row)) for row in rows])
        output_path.write_text(csv_content, encoding='utf-8')
        return output_path


def export_pptx(slides_data: List[dict], filename: str) -> Path:
    """
    Export data as PowerPoint PPTX

    Args:
        slides_data: List of dicts with 'title' and 'content' keys
        filename: Output filename

    Returns:
        Path to generated PPTX
    """
    output_path = DATA / "outputs" / filename

    if not PPTX_AVAILABLE:
        logging.warning("python-pptx not available, saving as TXT instead")
        output_path = DATA / "outputs" / filename.replace('.pptx', '.txt')
        txt_content = '\n\n'.join([f"Slide: {s['title']}\n{s['content']}" for s in slides_data])
        output_path.write_text(txt_content, encoding='utf-8')
        return output_path

    try:
        prs = Presentation()

        for slide_data in slides_data:
            # Add slide with title and content layout
            slide_layout = prs.slide_layouts[1]  # Title and Content
            slide = prs.slides.add_slide(slide_layout)

            # Add title
            title = slide.shapes.title
            title.text = slide_data.get('title', 'Untitled')

            # Add content
            content_placeholder = slide.shapes.placeholders[1]
            tf = content_placeholder.text_frame

            content = slide_data.get('content', '')
            for line in content.split('\n'):
                if line.strip():
                    p = tf.add_paragraph()
                    p.text = line.strip()
                    p.level = 0

        # Save
        prs.save(str(output_path))
        logging.info(f"Generated PPTX: {filename}")
        return output_path

    except Exception as e:
        logging.error(f"PPTX generation failed: {e}")
        # Fallback to text
        output_path = DATA / "outputs" / filename.replace('.pptx', '.txt')
        txt_content = '\n\n'.join([f"Slide: {s['title']}\n{s['content']}" for s in slides_data])
        output_path.write_text(txt_content, encoding='utf-8')
        return output_path


def extract_text(file_path: Union[str, Path]) -> str:
    """
    Extract text from various file formats

    Args:
        file_path: Path to file

    Returns:
        Extracted text content
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()

    try:
        # Text files
        if ext in ['.txt', '.md', '.py', '.js', '.java', '.cpp', '.c', '.html', '.css', '.json', '.xml']:
            return path.read_text(encoding='utf-8')

        # DOCX
        elif ext == '.docx' and DOCX_AVAILABLE:
            doc = Document(str(path))
            return '\n'.join([para.text for para in doc.paragraphs])

        # PDF
        elif ext == '.pdf' and PYPDF2_AVAILABLE:
            with open(path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                pages = []
                for page in reader.pages:
                    pages.append(page.extract_text())
                return '\n\n'.join(pages)

        # XLSX
        elif ext in ['.xlsx', '.xls'] and OPENPYXL_AVAILABLE:
            import pandas as pd
            df = pd.read_excel(path)
            return df.to_string()

        else:
            return f"Cannot extract text from {ext} file (missing library or unsupported format)"

    except Exception as e:
        logging.error(f"Text extraction failed for {path}: {e}")
        return f"Error extracting text: {str(e)}"
